import base64
import io
import logging
from datetime import datetime, date, timedelta
from typing import Dict, Any, List, Tuple

import xlrd

from odoo import models, fields, api, _
from odoo.exceptions import UserError

_logger = logging.getLogger(__name__)


class OpenpyxlSheetWrapper:
    def __init__(self, sheet):
        self._rows = list(sheet.iter_rows(values_only=True))

    @property
    def nrows(self):
        return len(self._rows)

    @property
    def ncols(self):
        return len(self._rows[0]) if self._rows else 0

    def cell_value(self, row, col):
        if row < len(self._rows) and col < len(self._rows[row]):
            val = self._rows[row][col]
            return val if val is not None else ''
        return ''


class OpenpyxlWorkbookWrapper:
    def __init__(self, wb):
        self._wb = wb

    def sheet_by_index(self, idx):
        return OpenpyxlSheetWrapper(self._wb.worksheets[idx])

    @property
    def datemode(self):
        return 0


class CSVWorkbookWrapper:
    def __init__(self, rows):
        self._sheet = CSVSheetWrapper(rows)

    def sheet_by_index(self, idx):
        return self._sheet

    @property
    def datemode(self):
        return 0


class CSVSheetWrapper:
    def __init__(self, rows):
        self._rows = rows

    @property
    def nrows(self):
        return len(self._rows)

    @property
    def ncols(self):
        return len(self._rows[0]) if self._rows else 0

    def cell_value(self, row, col):
        if row < len(self._rows) and col < len(self._rows[row]):
            return self._rows[row][col]
        return ''


class AttendanceUploadWizard(models.TransientModel):
    _name = 'bm.attendance.upload.wizard'
    _description = 'Upload Biometric Attendance XLS'

    file = fields.Binary(string='XLS File', required=True)
    file_name = fields.Char(string='File Name')
    state = fields.Selection([('draft', 'Draft'), ('done', 'Done')], default='draft')
    total_rows = fields.Integer(string='Total Rows', readonly=True)
    created_pairs = fields.Integer(string='New Records Created', readonly=True)
    skipped_rows = fields.Integer(string='Skipped Rows', readonly=True)
    unmapped_employees = fields.Integer(string='Unmapped (* Rows)', readonly=True)
    employees_created = fields.Integer(string='Employees Auto-Created', readonly=True)
    log = fields.Text(string='Details', readonly=True)
    imported_attendance_ids = fields.Many2many('hr.attendance', string='All Records', readonly=True)

    def action_import(self) -> Dict[str, Any]:
        self.ensure_one()
        if not self.file:
            raise UserError(_("Please upload an XLS file."))

        workbook = self._open_workbook()
        sheet = workbook.sheet_by_index(0)
        if sheet.nrows < 2:
            raise UserError(_("The sheet is empty or has no data rows."))

        header = self._read_header(sheet)
        col_idx = self._detect_columns(header)
        records, stats = self._parse_rows(sheet, col_idx, workbook.datemode)
        created, all_ids = self._create_attendance_records(records)

        # Count existing (total shown - new = existing)
        existing_count = len(all_ids) - created

        stats['created_pairs'] = created
        detail_lines = [
            f"Total rows in file        : {stats['total']}",
            f"Records parsed            : {stats['parsed']}",
            f"New records created       : {created}",
            f"Already existed (skipped) : {existing_count}",
            f"Employees auto-created    : {stats['employees_created']}",
            f"Skipped (invalid)         : {stats['skipped']}",
            f"Unmapped (*) rows         : {stats['unmapped']}",
        ]
        if stats['unmapped_enrolls']:
            detail_lines.append("Unmapped Enrolls: " + ", ".join(sorted(stats['unmapped_enrolls'])))

        self.write({
            'total_rows': stats['total'],
            'created_pairs': created,
            'skipped_rows': stats['skipped'],
            'unmapped_employees': stats['unmapped'],
            'employees_created': stats['employees_created'],
            'log': "\n".join(detail_lines),
            'state': 'done',
            'imported_attendance_ids': [(6, 0, all_ids)],
        })

        return {
            'type': 'ir.actions.act_window',
            'name': _('Biometric Attendance - All Records'),
            'res_model': 'hr.attendance',
            'view_mode': 'list',
            'views': [(self.env.ref('ma_bm_attendance.hr_attendance_biometric_list').id, 'list')],
            'domain': [('id', 'in', all_ids)],
            'target': 'current',
        }

    def _get_file_bytes(self) -> bytes:
        return base64.b64decode(self.file)

    def _open_workbook(self):
        file_bytes = self._get_file_bytes()

        if file_bytes[:2] == b'PK':
            try:
                import openpyxl
                wb = openpyxl.load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
                return OpenpyxlWorkbookWrapper(wb)
            except Exception as e:
                raise UserError(_("Cannot read XLSX file: %s") % e)

        if file_bytes[:4] == b'\xd0\xcf\x11\xe0':
            try:
                return xlrd.open_workbook(file_contents=file_bytes)
            except Exception:
                try:
                    import openpyxl
                    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
                    return OpenpyxlWorkbookWrapper(wb)
                except Exception:
                    pass
            text = file_bytes.decode('utf-8', errors='ignore')
            return self._parse_csv_text(text)

        text = file_bytes.decode('utf-8', errors='ignore').strip()
        return self._parse_csv_text(text)

    def _parse_csv_text(self, text):
        import csv
        lines = text.splitlines()
        for delim in [',', '\t', ';', '|']:
            rows = list(csv.reader(lines, delimiter=delim))
            if len(rows) > 1 and len(rows[0]) >= 4:
                return CSVWorkbookWrapper(rows)
        raise UserError(_("Could not parse file. Please save as .xlsx and try again."))

    @staticmethod
    def _read_header(sheet) -> List[str]:
        return [str(sheet.cell_value(0, c)).strip().lower() for c in range(sheet.ncols)]

    @staticmethod
    def _detect_columns(header: List[str]) -> Dict[str, int]:
        mapping: Dict[str, int] = {}
        for i, h in enumerate(header):
            if h == 'head':
                mapping['head'] = i
            elif h == 'date':
                mapping['date'] = i
            elif h in ('time to', 'time_to'):
                mapping['time_to'] = i
            elif h in ('check in', 'check_in'):
                mapping['check_in'] = i
            elif h in ('enroll number', 'enrollnumber', 'enroll no', 'enroll_no'):
                mapping['enroll'] = i
            elif h == 'name':
                mapping['name'] = i
            elif h in ('device number', 'devicenumber', 'device no', 'device_no'):
                mapping['device'] = i
            elif h in ('check out', 'check_out'):
                mapping['check_out'] = i
            elif h in ('time from', 'time_from'):
                mapping['time_from'] = i

        for required in ('date', 'enroll', 'name'):
            if required not in mapping:
                raise UserError(
                    _("Required column '%s' not found.\nDetected headers: %s")
                    % (required, ', '.join(header))
                )
        return mapping

    @api.model
    def _get_or_create_employee(self, name: str, enroll: str) -> int:
        if not name or name.strip() == '*':
            return False

        name = name.strip()

        emp = self.env['hr.employee'].search([('name', '=ilike', name)], limit=1)
        if emp:
            if not emp.device_user_id and enroll:
                emp.device_user_id = enroll
            return emp.id

        if enroll:
            emp = self.env['hr.employee'].search([('device_user_id', '=', enroll)], limit=1)
            if emp:
                return emp.id

        emp = self.env['hr.employee'].create({
            'name': name,
            'device_user_id': enroll or False,
        })
        return emp.id

    def _parse_rows(self, sheet, col_idx: Dict[str, int], datemode: int) -> Tuple[List[Dict], Dict]:
        records = []
        stats = {
            'total': 0, 'parsed': 0, 'skipped': 0, 'unmapped': 0,
            'unmapped_enrolls': set(), 'employees_created': 0, 'created_employees': set(),
        }
        name_cache = {}

        for r in range(1, sheet.nrows):
            stats['total'] += 1
            try:
                head = str(sheet.cell_value(r, col_idx.get('head', 0))).strip() if 'head' in col_idx else ''
                
                date_val = sheet.cell_value(r, col_idx['date'])
                attendance_date = self._parse_date(date_val, datemode)
                
                time_to = ''
                if 'time_to' in col_idx:
                    time_to_val = sheet.cell_value(r, col_idx['time_to'])
                    time_to = self._parse_time(time_to_val)
                
                check_in_code = ''
                if 'check_in' in col_idx:
                    check_in_val = sheet.cell_value(r, col_idx['check_in'])
                    check_in_code = self._parse_in_out_code(check_in_val)
                
                enroll_raw = sheet.cell_value(r, col_idx['enroll'])
                if isinstance(enroll_raw, (int, float)):
                    enroll = str(int(enroll_raw)).zfill(10)
                else:
                    enroll = str(enroll_raw).strip().zfill(10)
                
                name = str(sheet.cell_value(r, col_idx['name'])).strip()
                
                device = ''
                if 'device' in col_idx:
                    dev_raw = sheet.cell_value(r, col_idx['device'])
                    try:
                        if isinstance(dev_raw, (int, float)):
                            device = str(int(dev_raw)).zfill(4)
                        else:
                            device = str(int(float(str(dev_raw)))).zfill(4)
                    except:
                        device = str(dev_raw).strip()
                
                check_out_code = ''
                if 'check_out' in col_idx:
                    check_out_val = sheet.cell_value(r, col_idx['check_out'])
                    check_out_code = self._parse_in_out_code(check_out_val)
                
                time_from = ''
                if 'time_from' in col_idx:
                    time_from_val = sheet.cell_value(r, col_idx['time_from'])
                    time_from = self._parse_time(time_from_val)

            except Exception as e:
                _logger.warning("Row %s: parse error: %s", r, e)
                stats['skipped'] += 1
                continue

            if not name or name.strip() == '*':
                stats['unmapped'] += 1
                stats['unmapped_enrolls'].add(enroll)
                continue

            cache_key = name.lower().strip()
            if cache_key not in name_cache:
                emp_id = self._get_or_create_employee(name, enroll)
                name_cache[cache_key] = emp_id
            emp_id = name_cache[cache_key]

            if not emp_id:
                stats['unmapped'] += 1
                stats['unmapped_enrolls'].add(enroll)
                continue

            check_in_dt = False
            check_out_dt = False
            if time_to and attendance_date:
                try:
                    check_in_dt = datetime.combine(attendance_date, datetime.strptime(time_to, '%H:%M:%S').time())
                except:
                    pass
            if time_from and attendance_date:
                try:
                    check_out_dt = datetime.combine(attendance_date, datetime.strptime(time_from, '%H:%M:%S').time())
                except:
                    pass

            records.append({
                'employee_id': emp_id,
                'head': head,
                'attendance_date': attendance_date,
                'time_to': time_to,
                'check_in_code': check_in_code,
                'enroll_number': enroll,
                'employee_name': name,
                'device_number': device,
                'check_out_code': check_out_code,
                'time_from': time_from,
                'check_in_dt': check_in_dt,
                'check_out_dt': check_out_dt,
            })
            stats['parsed'] += 1

        return records, stats

    def _parse_date(self, val, datemode) -> date:
        if isinstance(val, datetime):
            return val.date()
        if isinstance(val, date):
            return val
        if isinstance(val, (int, float)):
            try:
                dt = xlrd.xldate.xldate_as_datetime(float(val), datemode)
                return dt.date()
            except:
                pass
        
        date_str = str(val).strip()
        formats = [
            '%d-%m-%Y',
            '%Y-%m-%d',
            '%d/%m/%Y',
            '%Y/%m/%d',
            '%m-%d-%Y',
            '%d.%m.%Y',
        ]
        for fmt in formats:
            try:
                return datetime.strptime(date_str, fmt).date()
            except ValueError:
                continue
        
        raise ValueError(f"Cannot parse date: {date_str}")

    def _parse_time(self, val) -> str:
        if isinstance(val, datetime):
            return val.strftime('%H:%M:%S')
        if isinstance(val, (int, float)):
            if val < 1:
                total_seconds = val * 86400
            else:
                total_seconds = val
            hours = int(total_seconds // 3600)
            minutes = int((total_seconds % 3600) // 60)
            seconds = int(total_seconds % 60)
            return f"{hours:02d}:{minutes:02d}:{seconds:02d}"
        return str(val).strip()

    def _parse_in_out_code(self, val) -> str:
        if isinstance(val, (int, float)):
            return str(int(val)).zfill(4)
        try:
            return str(int(float(str(val)))).zfill(4)
        except:
            return str(val).strip()

    @api.model
    def _create_attendance_records(self, records: List[Dict]) -> Tuple[int, List[int]]:
        if not records:
            return 0, []

        HrAttendance = self.env['hr.attendance']
        imported_ids = []
        existing_ids = []
        all_record_ids = []

        for rec in records:
            existing = HrAttendance.search([
                ('employee_id', '=', rec['employee_id']),
                ('attendance_date', '=', rec['attendance_date']),
                ('time_to', '=', rec['time_to']),
                ('is_biometric', '=', True),
            ], limit=1)

            if existing:
                existing_ids.append(existing.id)
                all_record_ids.append(existing.id)
            else:
                check_in_dt = rec['check_in_dt']
                check_out_dt = rec['check_out_dt']
                
                if check_in_dt and check_out_dt and check_out_dt <= check_in_dt:
                    check_out_dt = check_out_dt + timedelta(days=1)
                
                native_check_in = False
                native_check_out = False
                if check_in_dt and check_out_dt and check_out_dt > check_in_dt:
                    native_check_in = check_in_dt
                    native_check_out = check_out_dt
                elif check_in_dt:
                    native_check_in = check_in_dt

                record = HrAttendance.create({
                    'employee_id': rec['employee_id'],
                    'head': rec['head'],
                    'attendance_date': rec['attendance_date'],
                    'time_to': rec['time_to'],
                    'check_in_code': rec['check_in_code'],
                    'enroll_number': rec['enroll_number'],
                    'device_number': rec['device_number'],
                    'check_out_code': rec['check_out_code'],
                    'time_from': rec['time_from'],
                    'check_in': native_check_in,
                    'check_out': native_check_out,
                    'is_biometric': True,
                })
                imported_ids.append(record.id)
                all_record_ids.append(record.id)

        return len(imported_ids), all_record_ids